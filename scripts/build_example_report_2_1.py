#!/usr/bin/env python3
"""Build the Session 2.1 example report from Report_template.docx.

Output: examples/James_S1234567_W5_Menu.docx, a completed student report that
shows what a good report looks like. Prompts are read from 1169.md so they
match the session page exactly. Requires python-docx.
"""

from __future__ import annotations

import copy
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "Report_template.docx"
OUT = ROOT / "examples" / "James_S1234567_W5_Menu.docx"


# ---------- source prompts from 1169.md ----------

def session_text() -> str:
    md = (ROOT / "1169.md").read_text(encoding="utf-8")
    return md[md.index("#### Session 2.1:"):md.index("#### Session 2.2:")]


def fence_after(section: str, marker: str) -> str:
    start = section.index(marker)
    m = re.search(r"```text\n([\s\S]*?)\n[ \t>]*```", section[start:])
    lines = [re.sub(r"^[ \t]*>? ?", "", l) for l in m.group(1).splitlines()]
    return "\n".join(lines)


# ---------- docx helpers ----------

def clear_body_after_name(doc: Document) -> None:
    body = doc.element.body
    paras = doc.paragraphs
    keep = {paras[0]._p, paras[1]._p, paras[2]._p}
    for child in list(body):
        if child.tag == qn("w:sectPr") or child in keep:
            continue
        body.remove(child)


def shade(paragraph, fill: str = "F2F2F2") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def para(doc, text: str = "", bold_label: str | None = None, indent_first=True, italic=False):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent_first:
        p.paragraph_format.first_line_indent = Pt(36)
    if bold_label:
        p.add_run(bold_label).bold = True
    r = p.add_run(text)
    r.italic = italic
    return p


def bullet(doc, text: str, label: str | None = None):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Pt(28)
    p.paragraph_format.first_line_indent = Pt(-14)
    p.paragraph_format.space_after = Pt(3)
    p.add_run("-  ")
    if label:
        p.add_run(label).bold = True
    p.add_run(text)
    return p


def prompt_box(doc, text: str) -> None:
    for line in text.splitlines() or [""]:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Pt(14)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        shade(p)
        r = p.add_run(line if line else " ")
        r.font.name = "Courier New"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        r.font.size = Pt(9.5)
    doc.add_paragraph(style="Normal").paragraph_format.space_after = Pt(2)


def ai_box(doc, text: str) -> None:
    """AI answer excerpt: shaded, regular font, marked as quoted."""
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Pt(14)
    shade(p, "EAF1FB")
    r = p.add_run(text)
    r.font.size = Pt(10.5)


def caption(doc, text: str) -> None:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(10.5)


def set_cell_border(cell) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    tcpr.append(borders)


NUMERIC = re.compile(r"^\$?[\d,]+(\.\d+)?%?$")


def format_cell(cell, text: str, bold: bool, align) -> None:
    """Middle-aligned cell with no inherited paragraph spacing."""
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)
    pf.first_line_indent = Pt(0)
    pf.line_spacing = 1.0
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(10)
    set_cell_border(cell)


def table(doc, head: list[str], rows: list[list[str]], highlight: int | None = None) -> None:
    """Header and text columns centered; numeric columns right; all cells vertically centered."""
    t = doc.add_table(rows=1 + len(rows), cols=len(head))
    numeric_cols = {i for i in range(len(head)) if all(NUMERIC.match(r[i]) for r in rows)}
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        format_cell(c, h, True, WD_ALIGN_PARAGRAPH.CENTER)
        tcpr = c._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "D9D9D9")
        tcpr.append(shd)
    for r_i, row in enumerate(rows, start=1):
        for c_i, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.RIGHT if c_i in numeric_cols else WD_ALIGN_PARAGRAPH.CENTER
            format_cell(t.rows[r_i].cells[c_i], value, highlight is not None and r_i - 1 == highlight, align)
    doc.add_paragraph(style="Normal").paragraph_format.space_after = Pt(2)


def heading(doc, text: str) -> None:
    doc.add_paragraph(text, style="Heading 1")


# ---------- report content ----------

def build() -> None:
    sec = session_text()
    p_poor1 = fence_after(sec, "**Poor prompt 1**")
    p_imp1 = fence_after(sec, "**Improved prompt 1")
    p_imp2 = fence_after(sec, "**Improved prompt 2")
    p_partA = fence_after(sec, "**Part A - Worked example")
    p_partB = (
        "Continue with the same sales data and the same definitions.\n"
        "One thing has changed: the price of Taro Smoothie rose from $6.00 to $7.00 per cup.\n"
        "Keep the milk cost increase from Part A.\n"
        "Recalculate gross profit per cup and weekly gross profit, and show the\n"
        "new table sorted by weekly gross profit, from highest to lowest.\n"
        "Show the calculation for the Taro Smoothie row so I can check it.\n"
        "Then answer: should I still remove Taro Smoothie?\n"
        "Separate facts and guesses, and give a one-sentence recommendation."
    )

    doc = Document(TEMPLATE)
    clear_body_after_name(doc)
    paras = doc.paragraphs
    for p, text in zip(paras[:3], ["Session 2.1: Read Business Data Before Deciding", "S1234567", "James"]):
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
        p.runs[0].text = text

    note = doc.add_paragraph(style="Normal")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = note.add_run("Example report for students. Your AI answers may use different words.")
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    # Introduction
    heading(doc, "Introduction")
    para(doc, "Mina runs a small bubble tea shop near a university. At lunch, many students arrive at the "
              "same time, and the line is too long because staff must prepare five different drinks. Mina "
              "wants to remove one drink from the menu to make service faster, but she does not want to lose "
              "much profit. When she asked ChatGPT which drink to remove, it chose the drink with the fewest "
              "sales. My goal was to write prompts that help the AI find the drink that loses the least gross "
              "profit, and to check whether the AI answers are correct.")

    # Method
    heading(doc, "Method")
    para(doc, "I used five prompts in ChatGPT. I used a new chat for Prompt 1 and another new chat for "
              "Prompts 2 to 5, so the first answer could not affect the others.")
    steps = [
        ("Prompt 1 (Exercise 1, Step 1: Poor Prompt 1)", p_poor1,
         "I used this prompt to see what the AI does when the question has no decision rule."),
        ("Prompt 2 (Exercise 1, Step 2: Improved Prompt 1)", p_imp1,
         "I explained each column and gave the gross profit formula so the AI would read the data correctly "
         "and calculate the right measure."),
        ("Prompt 3 (Exercise 1, Step 4: Improved Prompt 2)", p_imp2,
         "I gave the AI the numbers I had already checked and my goal, so it would not recalculate and "
         "would separate facts from guesses."),
        ("Prompt 4 (Exercise 2, Part A: worked example)", p_partA,
         "I wanted to see whether my decision still holds when milk becomes more expensive."),
        ("Prompt 5 (Exercise 2, Part B: my own prompt from the template)", p_partB,
         "I filled in the template to test my friend's idea of raising the Taro Smoothie price instead of "
         "removing it."),
    ]
    for label, text, why in steps:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.space_before = Pt(6)
        p.add_run(label).bold = True
        prompt_box(doc, text)
        para(doc, why, bold_label="Why: ", indent_first=False)

    # Results
    heading(doc, "Results")
    caption(doc, "Result of Prompt 1 (poor prompt)")
    ai_box(doc, 'AI answer (key part): "Mango Fruit Tea has the lowest sales, with only 80 cups this week. '
                'I recommend removing it to simplify the menu."')
    para(doc, "The AI used cups sold as the decision rule because my question did not give a rule. "
              "Table 1 shows that this answer was wrong for Mina's goal.")

    caption(doc, "Table 1: Weekly gross profit by drink (AI answer to Prompt 2, checked by me)")
    table(doc, ["Drink", "Cups sold", "Revenue", "Gross profit per cup", "Weekly gross profit"], [
        ["Classic Milk Tea", "380", "$1,710.00", "$2.90", "$1,102.00"],
        ["Brown Sugar Milk Tea", "420", "$2,100.00", "$2.60", "$1,092.00"],
        ["Matcha Latte", "160", "$880.00", "$3.40", "$544.00"],
        ["Mango Fruit Tea", "80", "$440.00", "$4.20", "$336.00"],
        ["Taro Smoothie", "110", "$660.00", "$1.20", "$132.00"],
    ], highlight=4)
    para(doc, "My check: I checked the Taro Smoothie row with a calculator: $6.00 - $4.80 = $1.20 per cup, "
              "and 110 x $1.20 = $132.00. I also checked the Mango Fruit Tea row: $5.50 - $1.30 = $4.20 per "
              "cup, and 80 x $4.20 = $336.00. Both match the AI table, and the other three rows match the "
              "answer key.", indent_first=False)

    caption(doc, "Result of Prompt 3 (improved prompt)")
    ai_box(doc, 'AI answer (key part): "Facts: Taro Smoothie has the lowest weekly gross profit, $132.00. '
                'Mango Fruit Tea sells the fewest cups but earns $336.00 per week. Guesses: The data does not '
                'show how long each drink takes to make. Recommendation: Remove Taro Smoothie, because it loses '
                'the least gross profit. It is a blended drink that takes longer to make, so removing it will '
                'also shorten the line."')
    para(doc, "AI error I found: the AI correctly said under Guesses that the data does not show preparation "
              "time. But in the recommendation it wrote that Taro Smoothie \"takes longer to make\" as if it "
              "were a fact. This is only a guess, so I did not use it as a reason in my conclusion.",
         indent_first=False)

    caption(doc, "Table 2: Poor Prompt 1 compared with Improved Prompt 2")
    table(doc, ["Prompt", "Drink chosen", "Rule the AI used"], [
        ["Poor Prompt 1", "Mango Fruit Tea", "Fewest cups sold"],
        ["Improved Prompt 2", "Taro Smoothie", "Least weekly gross profit lost"],
    ])
    para(doc, "The answers are different because the improved prompts told the AI how to calculate gross "
              "profit and what my goal was, so it used gross profit instead of sales.", indent_first=False)

    caption(doc, "Table 3: Weekly gross profit after milk costs rise by $0.50 (Prompt 4, Part A)")
    table(doc, ["Drink", "Gross profit per cup", "Weekly gross profit"], [
        ["Classic Milk Tea", "$2.40", "$912.00"],
        ["Brown Sugar Milk Tea", "$2.10", "$882.00"],
        ["Matcha Latte", "$2.90", "$464.00"],
        ["Mango Fruit Tea", "$4.20", "$336.00"],
        ["Taro Smoothie", "$0.70", "$77.00"],
    ], highlight=4)
    para(doc, "My check: $6.00 - $5.30 = $0.70 per cup, and 110 x $0.70 = $77.00. Taro Smoothie is still the "
              "lowest. Total weekly gross profit fell from $3,206.00 to $2,671.00, a drop of $535.00.",
         indent_first=False)

    caption(doc, "Table 4: Taro Smoothie at $7.00, with the milk cost increase (Prompt 5, Part B)")
    table(doc, ["Drink", "Gross profit per cup", "Weekly gross profit"], [
        ["Classic Milk Tea", "$2.40", "$912.00"],
        ["Brown Sugar Milk Tea", "$2.10", "$882.00"],
        ["Matcha Latte", "$2.90", "$464.00"],
        ["Mango Fruit Tea", "$4.20", "$336.00"],
        ["Taro Smoothie", "$1.70", "$187.00"],
    ], highlight=4)
    para(doc, "My check: $7.00 - $5.30 = $1.70 per cup, and 110 x $1.70 = $187.00. The AI also listed a "
              "guess: the data does not show whether 110 customers will still buy Taro Smoothie at the "
              "higher price.", indent_first=False)

    # Conclusion
    heading(doc, "Conclusion")
    para(doc, "Mina should remove Taro Smoothie from the lunch menu. Even if its price rises to $7.00, it has "
              "the lowest weekly gross profit, $187.00, so removing it loses the least money. However, the "
              "data does not show how long each drink takes to make, so I cannot be sure that this change "
              "will shorten the line the most. The data also does not show whether customers would still "
              "buy Taro Smoothie at a higher price. Next time, I will tell the AI my decision rule and "
              "definitions before I ask for a recommendation.")

    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    build()
